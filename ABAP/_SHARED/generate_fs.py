# -*- coding: utf-8 -*-
"""
===============================================================================
 GENERATOR FUNCTIONAL SPECIFICATION - PROGRAM SAP
===============================================================================
 Cara pakai:
   1. Ubah bagian KONFIGURASI di bawah sesuai kondisi program terkini.
   2. Jalankan:  python generate_fs.py
   3. Hasil:     FS_<nama program>.docx  (beserta gambar diagram .png)

 Kebutuhan (sekali saja):
   pip install python-docx cairosvg

 Catatan: skrip ini TIDAK terhubung ke SAP. Isi data di bawah diketik manual
 dari hasil pengecekan di sistem (SE38 / SE16 tabel ZMAP_TYPE / SM36).
===============================================================================
"""

import os

# =============================================================================
# ============================ KONFIGURASI ====================================
# Bagian ini yang perlu diubah bila program / mapping berubah.
# =============================================================================

DOK = {
    "judul":        "Functional Specification",
    "subjudul":     "Notifikasi Persetujuan Purchase Order — Sprint 2",
    "baris_info":   "Program ZMMI_PO_EMAIL · Sistem SAP ECC · Disusun oleh Tim ABAP",
    "nama_program": "ZMMI_PO_EMAIL",
    "file_keluar":  "FS_ZMMI_PO_EMAIL_Sprint2.docx",
}

RUANG_LINGKUP = {
    "pengantar": (
        "Program ZMMI_PO_EMAIL mengirimkan notifikasi kepada approver atas Purchase Order "
        "yang masih menunggu persetujuan, lengkap dengan penandaan (flag) atas PO yang "
        "memerlukan perhatian khusus."
    ),
    "termasuk": (
        "penentuan tingkat persetujuan berdasarkan nilai PO, penentuan approver, "
        "penandaan kondisi khusus (exception), serta pengiriman email beserta lampiran rincian."
    ),
    "tidak_termasuk": (
        "pelepasan otomatis PO bernilai kecil (sudah berjalan pada Sprint 1 melalui program "
        "ZMMI_PO_RELEASE), dan persetujuan melalui email. Persetujuan tetap dilakukan di dalam "
        "SAP pada transaksi ME29N."
    ),
}

# --- 2. Parameter layar seleksi ---------------------------------------------
PARAMETER_JUDUL = ["Nama", "Deskripsi", "Tipe", "Keterangan penggunaan"]
PARAMETER_LEBAR = [1.05, 1.35, 1.10, 3.30]          # satuan inci
PARAMETER_DATA = [
    ("S_BUKRS", "Company Code", "Select-option",
     "Batasi PO pada company code tertentu. Kosong = seluruh company code."),
    ("S_EKORG", "Purchasing Organization", "Select-option",
     "Batasi pada purchasing organization tertentu. Kosong = seluruh org."),
    ("S_EKGRP", "Purchasing Group", "Select-option",
     "Batasi pada purchasing group tertentu. Nilai ini juga dipakai untuk menentukan approver Tier T1."),
    ("S_BSART", "Document Type PO", "Select-option",
     "Batasi jenis PO. Dipakai bila ada tipe PO yang tidak perlu dinotifikasi."),
    ("S_EBELN", "Nomor PO", "Select-option",
     "Menjalankan untuk satu atau beberapa nomor PO. Umumnya dipakai saat pengujian."),
    ("P_TEST", "Test Run", "Checkbox (default aktif)",
     "Aktif: seluruh proses berjalan tetapi email TIDAK dikirim; daftar penerima ditampilkan "
     "di layar. Nonaktif: email dikirim."),
]
PARAMETER_SOROT = ["P_TEST"]     # baris yang diberi latar oranye

# --- 3. Variant --------------------------------------------------------------
VARIANT_JUDUL = ["Nama Variant", "Dibuat oleh", "Kegunaan"]
VARIANT_LEBAR = [1.45, 1.45, 3.90]
VARIANT_DATA = [
    ("AUTO EMAIL", "TRSTDEV",
     "Variant yang dipakai untuk penjadwalan background job. Berisi kombinasi filter yang "
     "disepakati serta pengaturan Test Run. Sebelum dijadwalkan produktif, pastikan Test Run "
     "pada variant ini dalam keadaan NONAKTIF."),
]

# --- 4.1 THRESHOLD -----------------------------------------------------------
THRESHOLD_KOLOM_JUDUL = ["Kolom", "Isi", "Contoh", "Arti"]
THRESHOLD_KOLOM_LEBAR = [1.05, 1.20, 1.05, 3.50]
THRESHOLD_KOLOM_DATA = [
    ("TYPE",  "THRESHOLD",             "THRESHOLD",
     "Menandai baris ini sebagai konfigurasi ambang nilai."),
    ("OPT",   "Kode tier",             "T0 / T1 / T2",
     "Nama tier yang muncul di layar hasil dan email."),
    ("VALUE", "Batas atas nilai (USD)", "500.00",
     "PO masuk tier ini bila nilainya kurang dari atau sama dengan angka ini. "
     "Tier diuji berurutan dari nilai terkecil."),
    ("TEXT1", "Penanda tier auto",     "X",
     "Bila diisi X, PO pada tier tersebut DILEWATI oleh program ini "
     "(sudah ditangani auto-release Sprint 1)."),
]

THRESHOLD_DATA_JUDUL = ["Tier", "Batas nilai (USD)", "Penanda auto", "Perlakuan"]
THRESHOLD_DATA_LEBAR = [0.85, 1.55, 1.05, 3.35]
# (tier, batas, penanda, perlakuan, warna)  warna: "hijau" / "" / "abu"
THRESHOLD_DATA = [
    ("T0", "sampai 500.00", "X",
     "Dilewati. Auto-release ditangani program Sprint 1.", "hijau"),
    ("T1", "di atas 500.00 s.d. 5.000,00", "(kosong)",
     "Dinotifikasi ke approver sesuai Purchasing Group.", ""),
    ("T2", "di atas 5.000,00", "(kosong)",
     "Dinotifikasi ke approver GM (baris GM_FIXED).", "abu"),
]

# --- 4.2 APPROVER ------------------------------------------------------------
APPROVER_JUDUL = ["Kolom", "Isi", "Contoh", "Arti"]
APPROVER_LEBAR = [1.05, 1.45, 1.20, 3.10]
APPROVER_DATA = [
    ("TYPE",  "APPROVER", "APPROVER",
     "Menandai baris ini sebagai data approver."),
    ("OPT",   "Kunci grup approver", "P14 / GM_FIXED / DEFAULT",
     "Purchasing Group untuk Tier T1; GM_FIXED khusus Tier T2; DEFAULT sebagai cadangan "
     "bila grup tidak ditemukan."),
    ("VALUE", "Nama approver", "FENNY",
     "Ditampilkan pada sapaan email dan layar hasil. Ditulis apa adanya sesuai isi mapping."),
    ("TEXT1", "Nama pengguna email", "fenny", "Bagian sebelum tanda @."),
    ("TEXT2", "Domain email", "@trst.co.id",
     "Digabung dengan TEXT1 menjadi alamat email lengkap."),
    ("TEXT3", "Jenis kelamin", "M / F / kosong",
     "M menghasilkan sapaan Bapak, F menghasilkan Ibu, kosong hanya menyebut nama."),
]
APPROVER_CATATAN = (
    "Satu grup dapat memiliki lebih dari satu approver. Caranya cukup menambahkan baris baru "
    "dengan OPT yang sama dan nama yang berbeda. Seluruh alamat pada grup tersebut akan "
    "menerima email yang sama."
)

# --- 4.3 EXC_RULE ------------------------------------------------------------
EXCRULE_JUDUL = ["OPT", "VALUE", "TEXT1", "Arti aturan"]
EXCRULE_LEBAR = [1.35, 1.15, 0.95, 3.35]
EXCRULE_DATA = [
    ("VENDOR_NEW", "DAYS", "90",
     "Vendor dianggap baru bila tanggal pembuatan master kurang dari 90 hari."),
    ("VENDOR_NEW", "NO_HISTORY", "X",
     "Vendor dianggap baru bila belum pernah memiliki PO yang sudah dirilis."),
    ("MATERIAL_NEW", "DAYS", "90",
     "Material dianggap baru bila tanggal pembuatan master kurang dari 90 hari."),
    ("MATERIAL_NEW", "NO_HISTORY", "X",
     "Material dianggap baru bila belum pernah dibeli (tidak ada riwayat harga)."),
    ("PRICE_INCREASE", "PERCENT", "10",
     "Batas kenaikan harga. Ditandai bila harga naik lebih dari 10% dibanding harga terakhir."),
    ("PRICE_HISTORY", "STALE_DAYS", "730",
     "Riwayat harga dianggap usang bila pembelian terakhir lebih lama dari 730 hari (24 bulan)."),
    ("PSTYP", "9", "SvcCapex",
     "Item dengan kategori PSTYP = 9 (Service) ditandai dengan flag SvcCapex."),
    ("KNTTP", "A", "SvcCapex",
     "Item dengan account assignment KNTTP = A (Asset/Capex) ditandai dengan flag SvcCapex."),
]

# --- 5. Ketentuan approver & flagging ---------------------------------------
KETENTUAN_APPROVER = [
    "PO Tier T2 diarahkan ke baris mapping GM_FIXED.",
    "PO Tier T1 diarahkan ke baris mapping yang sesuai dengan Purchasing Group PO tersebut.",
    "Apabila Purchasing Group tidak ditemukan pada mapping, digunakan baris DEFAULT sebagai cadangan.",
    "Apabila baris DEFAULT juga tidak ada, PO tetap muncul pada layar hasil dengan status "
    "kesalahan dan tidak dikirimkan email.",
]

FLAG_JUDUL = ["Flag", "Kondisi ditandai", "Catatan penting"]
FLAG_LEBAR = [1.30, 2.35, 3.15]
FLAG_DATA = [
    ("Vendor Baru",
     "Umur master vendor kurang dari batas DAYS, ATAU vendor belum pernah punya PO yang dirilis.",
     "Dua aturan berdiri sendiri. Bila salah satu baris mapping dihapus, pengecekan itu saja "
     "yang berhenti."),
    ("Material Baru",
     "Umur master material kurang dari batas DAYS, ATAU material belum pernah dibeli.",
     "Diperiksa per baris item PO yang memiliki nomor material."),
    ("Harga Naik (Material)",
     "Harga SATUAN sekarang lebih tinggi dari batas persen dibanding harga terakhir material "
     "tersebut, dari vendor manapun.",
     "Perbandingan di level line item, bukan level PO. Harga dinormalisasi ke satu satuan "
     "(NETPR dibagi PEINH) lalu dikonversi ke USD."),
    ("Harga Naik (Vendor)",
     "Harga SATUAN sekarang lebih tinggi dari batas persen dibanding harga terakhir dari "
     "vendor yang sama untuk material tersebut.",
     "Diperiksa terpisah dari sumbu material, sehingga keduanya dapat muncul bersamaan."),
    ("Histori Usang",
     "Pembelian terakhir material tersebut lebih lama dari batas STALE_DAYS.",
     "Bila kondisi ini terpenuhi, perbandingan kenaikan harga dilewati karena acuannya "
     "dianggap tidak lagi relevan."),
    ("Jasa/Capex",
     "Item PO memiliki kategori PSTYP atau account assignment KNTTP yang terdaftar pada mapping.",
     "Daftar kategori dapat ditambah sewaktu-waktu dengan menambah baris pada mapping."),
]

CATATAN_HARGA = [
    "Perbandingan dilewati bila satuan harga (BPRME) berbeda antara PO sekarang dan PO acuan, "
    "misalnya sebelumnya per KG sekarang per TON. Program memilih tidak membandingkan daripada "
    "menghasilkan angka yang menyesatkan.",
    "Perbandingan juga dilewati bila riwayat harga sudah melewati batas STALE_DAYS. PO tersebut "
    "ditandai Histori Usang, bukan Harga Naik.",
]

# --- 6. Contoh email ---------------------------------------------------------
EMAIL_CONTOH = {
    "dari":    "TRST-BUDI@TRST.CO.ID",
    "kepada":  "budi.purwanto@trst.co.id ; nafianta.bp@trias-sentosa.com",
    "subjek":  "PO Menunggu Persetujuan - 30 Juli 2026",
    "lampiran": "Detail PO - 30 Juli 2026.XLS",
    "salam":   "Yth. Bapak GM Purchasing,",
    "kalimat": "Berikut Purchase Order yang menunggu persetujuan Anda (grup GM_FIXED):",
    "kolom":   ["No PO", "Vendor", "Nilai (USD)", "Tier", "Flag", "Alasan Exception"],
    # (no po, vendor, nilai, tier, flag, alasan, perhatian?)
    "baris": [
        ("4503000337", "2131000195", "50,000,000.00", "T2", "Perhatian",
         "Material Baru; Harga Naik (Vendor)", True),
        ("4508000108", "2131000732", "1,034,202.56", "T2", "Perhatian",
         "Histori Usang", True),
        ("4507000209", "2131000441", "10,000.00", "T2", "OK", "-", False),
    ],
    "legenda": [
        "Vendor Baru = vendor dibuat < 90 hari atau belum pernah bertransaksi.",
        "Material Baru = material dibuat < 90 hari atau belum pernah dibeli.",
        "Harga Naik (Material / Vendor) = harga naik di atas 10%.",
        "Histori Usang = pembelian terakhir lebih lama dari 24 bulan (730 hari).",
        "Jasa/Capex = kategori PSTYP=9 (Service), KNTTP=A (Asset/Capex).",
    ],
    "penutup":  "Mohon lakukan persetujuan melalui transaksi ME29N.",
    "footnote": "Email ini dikirim otomatis oleh sistem. Mohon tidak membalas email ini.",
}

ISI_EMAIL_POIN = [
    "Subjek: PO Menunggu Persetujuan diikuti tanggal pengiriman.",
    "Sapaan menggunakan nama approver sesuai mapping, dengan sebutan Bapak atau Ibu "
    "mengikuti kolom jenis kelamin.",
    "Tabel rincian berisi kolom: No PO, Vendor, Nilai (USD), Tier, Flag, dan Alasan Exception.",
    "Keterangan flag disusun otomatis mengikuti aturan yang sedang aktif, sehingga angka yang "
    "tertulis selalu sama dengan konfigurasi.",
    "Lampiran Excel dengan nama Detail PO diikuti tanggal, berisi kolom yang sama dengan "
    "tabel pada email.",
]

# --- 7. Batasan --------------------------------------------------------------
BATASAN_JUDUL = ["Hal", "Kondisi saat ini"]
BATASAN_LEBAR = [2.05, 4.75]
# (hal, kondisi, "merah"/"abu"/"")
BATASAN_DATA = [
    ("Pengiriman ke domain luar",
     "Masih ditolak mail server dengan pesan 550 Authenticated Relay not permitted. Selain itu, "
     "satu alamat yang ditolak menyebabkan seluruh email tersebut gagal, termasuk penerima "
     "internal. Sedang dikoordinasikan dengan Tim Basis.", "merah"),
    ("Alamat pengirim",
     "Masih menggunakan identitas pengguna SAP yang menjalankan program. Diusulkan memakai "
     "akun notifikasi khusus.", "merah"),
    ("Pengulangan notifikasi",
     "PO yang belum disetujui akan dikirim ulang pada setiap kali job berjalan. Belum ada "
     "pencatatan riwayat pengiriman.", "merah"),
    ("Persetujuan di dalam SAP",
     "Program hanya memberi notifikasi. Persetujuan tetap dilakukan pada transaksi ME29N, "
     "sesuai kesepakatan ruang lingkup.", ""),
    ("Pemeriksaan otorisasi",
     "Program belum melakukan pemeriksaan otorisasi organisasi. Dijalankan melalui background "
     "job dengan pengguna teknis.", "abu"),
]

# --- Langkah alur (diagram) --------------------------------------------------
# (label kecil, judul, baris penjelas 1, baris penjelas 2)
ALUR_ATAS = [
    ("LANGKAH 1", "Baca konfigurasi", "ZMAP_TYPE: THRESHOLD,", "APPROVER, EXC_RULE"),
    ("LANGKAH 2", "Ambil PO pending", "EKKO: FRGRL='X', BSTYP='F',", "tidak dihapus + filter layar"),
    ("LANGKAH 3", "Hitung nilai PO", "SUM EKPO-NETWR ->", "konversi USD kurs tgl PO"),
    ("LANGKAH 4", "Tentukan Tier", "banding ambang THRESHOLD", "(urut menaik)"),
]
ALUR_BAWAH = [
    ("LANGKAH 5", "Tentukan approver", "T2 -> GM_FIXED; T1 -> EKGRP;", "bila tak ada -> DEFAULT"),
    ("LANGKAH 6", "Exception flagging", "cek tiap aturan aktif", "(hanya Tier T1/T2)"),
    ("LANGKAH 7", "Kelompokkan per grup", "satu grup approver =", "satu email"),
    ("LANGKAH 8", "Kirim email", "HTML + lampiran Excel", "ke semua approver grup"),
]
ALUR_SKIP = ("TIER AUTO (T0)", "PO di-SKIP", "ditangani Sprint 1", "(ZMMI_PO_RELEASE)")
ALUR_TESTRUN = ("Parameter P_TEST (Test Run)",
                "Aktif = seluruh langkah 1-7 tetap berjalan, tetapi email TIDAK dikirim. "
                "Daftar penerima ditampilkan di layar hasil untuk verifikasi.")
ALUR_HASIL = ("Hasil akhir",
              "Email notifikasi per grup approver + daftar hasil di layar "
              "(nomor PO, nilai USD, tier, approver, status, flag exception).")

# =============================================================================
# ======================= AKHIR BAGIAN KONFIGURASI ============================
# =============================================================================


# ----------------------------- warna & huruf ---------------------------------
HEAD, ACC, MUT = "2F3A46", "C8622D", "6B7280"
LIGHT, ACCBG = "F4F5F7", "FBF1EA"
RED, REDBG = "B03636", "FBECEC"
GRN, GRNBG = "2E7D32", "EDF6EE"
BLU, BLUBG = "2F5D8C", "EAF1F8"
BORD, WHITE, TXT = "C9CED4", "FFFFFF", "1F2328"
FONT_SVG = "DejaVu Sans, Arial, sans-serif"

WARNA_LATAR = {"merah": REDBG, "abu": LIGHT, "hijau": GRNBG, "sorot": ACCBG, "": None}
WARNA_TEKS = {"merah": RED, "abu": ACC, "hijau": GRN, "sorot": ACC, "": ACC}


# ======================== BAGIAN 1 : GAMBAR DIAGRAM ==========================
def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _kotak(x, y, w, h, label, judul, sub1="", sub2="",
           isi=LIGHT, garis=BORD, warna_judul=TXT, warna_label=ACC, ukuran=13.5):
    cx = x + w / 2
    s = (f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="10" fill="#{isi}" '
         f'stroke="#{garis}" stroke-width="1.6"/>')
    s += (f'<text x="{cx}" y="{y+19}" text-anchor="middle" font-family="{FONT_SVG}" '
          f'font-size="10.5" font-weight="800" fill="#{warna_label}">{_esc(label)}</text>')
    s += (f'<text x="{cx}" y="{y+38}" text-anchor="middle" font-family="{FONT_SVG}" '
          f'font-size="{ukuran}" font-weight="700" fill="#{warna_judul}">{_esc(judul)}</text>')
    if sub1:
        s += (f'<text x="{cx}" y="{y+55}" text-anchor="middle" font-family="{FONT_SVG}" '
              f'font-size="11" fill="#{MUT}">{_esc(sub1)}</text>')
    if sub2:
        s += (f'<text x="{cx}" y="{y+70}" text-anchor="middle" font-family="{FONT_SVG}" '
              f'font-size="11" fill="#{MUT}">{_esc(sub2)}</text>')
    return s


def _panah_kanan(x1, x2, y, warna=MUT):
    return (f'<line x1="{x1}" y1="{y}" x2="{x2-9}" y2="{y}" stroke="#{warna}" stroke-width="2"/>'
            f'<polygon points="{x2-10},{y-6} {x2-10},{y+6} {x2},{y}" fill="#{warna}"/>')


def _panah_bawah(x, y1, y2, warna=MUT, label="", warna_label=None):
    s = (f'<line x1="{x}" y1="{y1}" x2="{x}" y2="{y2-9}" stroke="#{warna}" stroke-width="2"/>'
         f'<polygon points="{x-6},{y2-10} {x+6},{y2-10} {x},{y2}" fill="#{warna}"/>')
    if label:
        wl = warna_label or warna
        s += (f'<text x="{x+8}" y="{(y1+y2)/2+4}" font-family="{FONT_SVG}" font-size="10" '
              f'font-weight="700" fill="#{wl}">{_esc(label)}</text>')
    return s


def gambar_alur(path_png):
    """Diagram alur eksekusi program."""
    W, H = 1080, 520
    bw, bh, r1 = 238, 80, 42
    xs = [20, 278, 536, 794]
    r2, r3 = 176, 290

    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" '
           f'font-family="{FONT_SVG}">', f'<rect width="{W}" height="{H}" fill="#{WHITE}"/>']
    svg.append(f'<text x="20" y="24" font-family="{FONT_SVG}" font-size="13" '
               f'font-weight="800" fill="#{HEAD}">Alur eksekusi program {_esc(DOK["nama_program"])}</text>')

    for i, (lb, jd, s1, s2) in enumerate(ALUR_ATAS):
        gaya = dict(isi=BLUBG, garis=BLU, warna_judul=BLU, warna_label=BLU) if i == 0 else {}
        svg.append(_kotak(xs[i], r1, bw, bh, lb, jd, s1, s2, **gaya))
    for i in range(3):
        svg.append(_panah_kanan(xs[i] + bw, xs[i + 1], r1 + bh / 2))

    lb, jd, s1, s2 = ALUR_SKIP
    svg.append(_kotak(xs[3], r2, bw, 72, lb, jd, s1, s2,
                      isi=GRNBG, garis=GRN, warna_judul=GRN, warna_label=GRN))
    svg.append(_panah_bawah(xs[3] + bw / 2, r1 + bh, r2, GRN,
                            label="nilai <= ambang auto", warna_label=GRN))

    for i, (lb, jd, s1, s2) in enumerate(ALUR_BAWAH):
        gaya = dict(isi=ACCBG, garis=ACC, warna_judul=ACC, warna_label=ACC) if i < 2 else {}
        svg.append(_kotak(xs[i], r3, bw, bh, lb, jd, s1, s2, **gaya))
    for i in range(3):
        svg.append(_panah_kanan(xs[i] + bw, xs[i + 1], r3 + bh / 2))

    bx = xs[3] + 46
    svg.append(f'<line x1="{bx}" y1="{r1+bh}" x2="{bx}" y2="{r3-26}" stroke="#{ACC}" stroke-width="2"/>')
    svg.append(f'<text x="{bx-8}" y="{r1+bh+34}" text-anchor="end" font-family="{FONT_SVG}" '
               f'font-size="10" font-weight="700" fill="#{ACC}">Tier T1 / T2</text>')
    svg.append(f'<line x1="{bx}" y1="{r3-26}" x2="{xs[0]+bw/2}" y2="{r3-26}" stroke="#{ACC}" stroke-width="2"/>')
    svg.append(_panah_bawah(xs[0] + bw / 2, r3 - 26, r3, ACC))

    ty = 402
    svg.append(f'<rect x="20" y="{ty}" width="{W-40}" height="46" rx="9" fill="#{ACCBG}" '
               f'stroke="#{ACC}" stroke-width="1.5" stroke-dasharray="5 4"/>')
    svg.append(f'<text x="34" y="{ty+19}" font-family="{FONT_SVG}" font-size="11.5" '
               f'font-weight="800" fill="#{ACC}">{_esc(ALUR_TESTRUN[0])}</text>')
    svg.append(f'<text x="34" y="{ty+36}" font-family="{FONT_SVG}" font-size="11.5" '
               f'fill="#{TXT}">{_esc(ALUR_TESTRUN[1])}</text>')

    ny = 462
    svg.append(f'<rect x="20" y="{ny}" width="{W-40}" height="44" rx="9" fill="#{LIGHT}" '
               f'stroke="#{BORD}" stroke-width="1.4"/>')
    svg.append(f'<text x="34" y="{ny+18}" font-family="{FONT_SVG}" font-size="11.5" '
               f'font-weight="800" fill="#{HEAD}">{_esc(ALUR_HASIL[0])}</text>')
    svg.append(f'<text x="34" y="{ny+35}" font-family="{FONT_SVG}" font-size="11.5" '
               f'fill="#{TXT}">{_esc(ALUR_HASIL[1])}</text>')
    svg.append('</svg>')

    import cairosvg
    cairosvg.svg2png(bytestring="".join(svg).encode(), write_to=path_png,
                     scale=2.0, background_color="white")


def gambar_email(path_png):
    """Ilustrasi tampilan email notifikasi."""
    E = EMAIL_CONTOH
    W, H = 1000, 470
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" '
           f'font-family="{FONT_SVG}">', f'<rect width="{W}" height="{H}" fill="#{WHITE}"/>']
    svg.append(f'<rect x="16" y="16" width="{W-32}" height="{H-32}" rx="8" fill="#{WHITE}" '
               f'stroke="#{BORD}" stroke-width="1.6"/>')
    svg.append(f'<rect x="16" y="16" width="{W-32}" height="74" rx="8" fill="#{LIGHT}" '
               f'stroke="#{BORD}" stroke-width="1.6"/>')
    svg.append(f'<rect x="16" y="76" width="{W-32}" height="14" fill="#{LIGHT}"/>')

    for i, (lab, val, tebal, warna) in enumerate([
            ("Dari   :", E["dari"], "700", TXT),
            ("Kepada :", E["kepada"], "400", TXT),
            ("Subjek :", E["subjek"], "700", ACC)]):
        y = 40 + i * 18
        svg.append(f'<text x="34" y="{y}" font-family="{FONT_SVG}" font-size="11.5" '
                   f'fill="#{MUT}">{_esc(lab)}</text>')
        svg.append(f'<text x="96" y="{y}" font-family="{FONT_SVG}" font-size="11.5" '
                   f'font-weight="{tebal}" fill="#{warna}">{_esc(val)}</text>')

    svg.append(f'<rect x="34" y="100" width="230" height="26" rx="5" fill="#{GRNBG}" '
               f'stroke="#{GRN}" stroke-width="1.2"/>')
    svg.append(f'<text x="46" y="118" font-family="{FONT_SVG}" font-size="11" '
               f'font-weight="700" fill="#{GRN}">{_esc(E["lampiran"])}</text>')
    svg.append(f'<text x="34" y="152" font-family="{FONT_SVG}" font-size="12" '
               f'fill="#{TXT}">{_esc(E["salam"])}</text>')
    svg.append(f'<text x="34" y="174" font-family="{FONT_SVG}" font-size="12" '
               f'fill="#{TXT}">{_esc(E["kalimat"])}</text>')

    tx, ty, rh = 34, 190, 26
    cw = [112, 120, 110, 60, 96, 340]
    svg.append(f'<rect x="{tx}" y="{ty}" width="{sum(cw)}" height="{rh}" fill="#{HEAD}"/>')
    cx = tx
    for i, h in enumerate(E["kolom"]):
        svg.append(f'<text x="{cx+7}" y="{ty+17}" font-family="{FONT_SVG}" font-size="10.5" '
                   f'font-weight="700" fill="#{WHITE}">{_esc(h)}</text>')
        cx += cw[i]

    ry = ty + rh
    for b in E["baris"]:
        perhatian = b[6]
        svg.append(f'<rect x="{tx}" y="{ry}" width="{sum(cw)}" height="{rh}" '
                   f'fill="#{REDBG if perhatian else WHITE}" stroke="#{BORD}" stroke-width="0.9"/>')
        cx = tx
        for i in range(6):
            warna, tebal = TXT, "400"
            if i == 4:
                warna = RED if perhatian else GRN
                tebal = "700"
            svg.append(f'<text x="{cx+7}" y="{ry+17}" font-family="{FONT_SVG}" font-size="10.5" '
                       f'font-weight="{tebal}" fill="#{warna}">{_esc(b[i])}</text>')
            cx += cw[i]
        ry += rh

    ky = ry + 18
    svg.append(f'<text x="34" y="{ky}" font-family="{FONT_SVG}" font-size="11" '
               f'font-weight="700" fill="#444444">Keterangan flag (mengikuti mapping aktif):</text>')
    for i, t in enumerate(E["legenda"]):
        svg.append(f'<circle cx="42" cy="{ky+16+i*17-4}" r="2.2" fill="#444444"/>')
        svg.append(f'<text x="52" y="{ky+16+i*17}" font-family="{FONT_SVG}" font-size="10.5" '
                   f'fill="#444444">{_esc(t)}</text>')

    fy = ky + 16 + len(E["legenda"]) * 17 + 16
    svg.append(f'<text x="34" y="{fy}" font-family="{FONT_SVG}" font-size="11.5" '
               f'fill="#{TXT}">{_esc(E["penutup"])}</text>')
    svg.append(f'<text x="34" y="{fy+18}" font-family="{FONT_SVG}" font-size="10" '
               f'fill="#{MUT}">{_esc(E["footnote"])}</text>')
    svg.append('</svg>')

    import cairosvg
    cairosvg.svg2png(bytestring="".join(svg).encode(), write_to=path_png,
                     scale=2.0, background_color="white")


# ========================= BAGIAN 2 : DOKUMEN WORD ===========================
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _rgb(hx):
    return RGBColor(int(hx[0:2], 16), int(hx[2:4], 16), int(hx[4:6], 16))


def _arsir(sel, warna_hex):
    if not warna_hex:
        return
    tcPr = sel._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), warna_hex)
    tcPr.append(shd)


def _tulis_sel(sel, teks, tebal=False, warna=TXT, ukuran=8.7):
    sel.text = ""
    par = sel.paragraphs[0]
    par.paragraph_format.space_after = Pt(1)
    r = par.add_run(str(teks))
    r.bold = tebal
    r.font.size = Pt(ukuran)
    r.font.color.rgb = _rgb(warna)


def _paragraf(doc, teks, ukuran=9.5, warna=TXT, tebal=False, miring=False, spasi=5):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(spasi)
    r = par.add_run(teks)
    r.bold, r.italic = tebal, miring
    r.font.size = Pt(ukuran)
    r.font.color.rgb = _rgb(warna)
    return par


def _paragraf_campur(doc, potongan, ukuran=9.5, spasi=5):
    """potongan = [(teks, tebal, warna), ...]"""
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(spasi)
    for teks, tebal, warna in potongan:
        r = par.add_run(teks)
        r.bold = tebal
        r.font.size = Pt(ukuran)
        r.font.color.rgb = _rgb(warna)
    return par


def _judul(doc, teks, tingkat=1):
    par = doc.add_heading(level=tingkat)
    par.paragraph_format.space_before = Pt(12 if tingkat == 1 else 9)
    par.paragraph_format.space_after = Pt(5)
    r = par.add_run(teks)
    r.font.size = Pt(14 if tingkat == 1 else 11.5)
    r.bold = True
    r.font.color.rgb = _rgb(HEAD)
    return par


def _poin(doc, teks, ukuran=9.5):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(teks)
    r.font.size = Pt(ukuran)
    return par


def _kunci_lebar(t, lebar):
    """Kunci lebar kolom agar tidak diubah otomatis oleh Word/LibreOffice."""
    tblPr = t._tbl.tblPr
    for lama in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(lama)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    for i, kol in enumerate(t.columns):
        kol.width = Inches(lebar[i])


def _tabel(doc, judul_kolom, lebar, baris, warna_baris=None, tebal_kolom1=True):
    """baris = list of tuple; warna_baris = list kode warna sejajar baris."""
    t = doc.add_table(rows=1, cols=len(judul_kolom))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _kunci_lebar(t, lebar)

    for i, jd in enumerate(judul_kolom):
        sel = t.rows[0].cells[i]
        sel.width = Inches(lebar[i])
        _tulis_sel(sel, jd, tebal=True, warna=WHITE)
        _arsir(sel, HEAD)

    for n, data in enumerate(baris):
        kode = (warna_baris[n] if warna_baris else ("abu" if n % 2 else ""))
        latar = WARNA_LATAR.get(kode)
        warna1 = WARNA_TEKS.get(kode, ACC)
        sel_baris = t.add_row().cells
        for i, nilai in enumerate(data):
            sel = sel_baris[i]
            sel.width = Inches(lebar[i])
            _tulis_sel(sel, nilai,
                       tebal=(i == 0 and tebal_kolom1),
                       warna=(warna1 if i == 0 else TXT))
            _arsir(sel, latar)

    _kunci_lebar(t, lebar)      # dipasang ulang setelah baris ditambahkan
    return t


def _garis(doc):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(7)
    p = par._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'), 'single')
    bt.set(qn('w:sz'), '6')
    bt.set(qn('w:color'), 'D9DDE2')
    bd.append(bt)
    p.append(bd)


def _gambar(doc, path_png, lebar_inci, keterangan):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(2)
    par.add_run().add_picture(path_png, width=Inches(lebar_inci))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(keterangan)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(MUT)


def buat_dokumen(png_alur, png_email, path_keluar):
    doc = Document()

    gaya = doc.styles["Normal"]
    gaya.font.name = "Calibri"
    gaya.font.size = Pt(9.5)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    sec.left_margin = sec.right_margin = Inches(0.8)

    # ---------------- Sampul ----------------
    _paragraf(doc, DOK["judul"], ukuran=16, warna=HEAD, tebal=True, spasi=2)
    _paragraf(doc, DOK["subjudul"], ukuran=11, warna=ACC, tebal=True, spasi=2)
    _paragraf(doc, DOK["baris_info"], ukuran=8.5, warna=MUT, spasi=3)
    _garis(doc)

    # ---------------- 1. Ruang lingkup ----------------
    _judul(doc, "1. Ruang Lingkup")
    _paragraf(doc, RUANG_LINGKUP["pengantar"])
    _paragraf_campur(doc, [("Termasuk dalam ruang lingkup: ", True, TXT),
                           (RUANG_LINGKUP["termasuk"], False, TXT)])
    _paragraf_campur(doc, [("Tidak termasuk: ", True, TXT),
                           (RUANG_LINGKUP["tidak_termasuk"], False, TXT)])

    # ---------------- 2. Parameter ----------------
    _judul(doc, "2. Parameter Program")
    _paragraf(doc, "Seluruh parameter berada pada layar seleksi dan bersifat opsional, "
                   "kecuali disebutkan lain. Parameter kosong berarti tidak ada pembatasan.")
    kode = ["sorot" if b[0] in PARAMETER_SOROT else ("abu" if i % 2 else "")
            for i, b in enumerate(PARAMETER_DATA)]
    _tabel(doc, PARAMETER_JUDUL, PARAMETER_LEBAR, PARAMETER_DATA, warna_baris=kode)
    _paragraf_campur(doc, [("Catatan: ", True, MUT),
                           ("Test Run aktif secara bawaan agar program aman dijalankan untuk "
                            "pengecekan tanpa risiko mengirim email kepada approver.", False, MUT)],
                     ukuran=8.5)

    # ---------------- 3. Variant ----------------
    _judul(doc, "3. Variant Program")
    _paragraf(doc, "Variant menyimpan kombinasi isian layar seleksi agar penjadwalan job "
                   "tidak perlu diisi ulang setiap kali.")
    _tabel(doc, VARIANT_JUDUL, VARIANT_LEBAR, VARIANT_DATA)

    doc.add_page_break()

    # ---------------- 4. Mapping ----------------
    _judul(doc, "4. Konfigurasi pada Tabel ZMAP_TYPE")
    _paragraf(doc, "Seluruh aturan bisnis disimpan pada tabel konfigurasi ZMAP_TYPE sehingga "
                   "dapat diubah tanpa mengubah program. Baris yang ditandai sebagai dihapus "
                   "(kolom Deletion) diabaikan oleh program.")
    _paragraf(doc, "Terdapat tiga jenis konfigurasi: THRESHOLD (ambang nilai), APPROVER "
                   "(data penerima), dan EXC_RULE (aturan penandaan).")

    _judul(doc, "4.1 THRESHOLD — Ambang Nilai dan Tingkat Persetujuan", 2)
    _tabel(doc, THRESHOLD_KOLOM_JUDUL, THRESHOLD_KOLOM_LEBAR, THRESHOLD_KOLOM_DATA)
    _paragraf(doc, "Konfigurasi yang berlaku saat ini:")
    _tabel(doc, THRESHOLD_DATA_JUDUL, THRESHOLD_DATA_LEBAR,
           [b[:4] for b in THRESHOLD_DATA], warna_baris=[b[4] for b in THRESHOLD_DATA])

    _judul(doc, "4.2 APPROVER — Data Penerima Notifikasi", 2)
    _tabel(doc, APPROVER_JUDUL, APPROVER_LEBAR, APPROVER_DATA)
    _paragraf(doc, APPROVER_CATATAN)

    doc.add_page_break()

    _judul(doc, "4.3 EXC_RULE — Aturan Penandaan Kondisi Khusus", 2)
    _paragraf(doc, "Aturan penandaan bersifat dinamis. Apabila suatu baris dihapus atau "
                   "ditandai Deletion, pemeriksaan tersebut otomatis berhenti dijalankan "
                   "tanpa perlu mengubah program.")
    _tabel(doc, EXCRULE_JUDUL, EXCRULE_LEBAR, EXCRULE_DATA)

    # ---------------- 5. Alur ----------------
    _judul(doc, "5. Alur Eksekusi Program")
    _gambar(doc, png_alur, 6.55,
            "Gambar 1 — Alur eksekusi program dari pembacaan konfigurasi sampai pengiriman email")

    _judul(doc, "Ketentuan penentuan approver", 2)
    for t in KETENTUAN_APPROVER:
        _poin(doc, t)

    _judul(doc, "Ketentuan penandaan kondisi khusus (exception)", 2)
    _paragraf(doc, "Penandaan hanya dihitung untuk PO Tier T1 dan T2. Satu PO dapat memiliki "
                   "lebih dari satu penandaan; seluruh alasan ditampilkan pada email dan lampiran.")
    _tabel(doc, FLAG_JUDUL, FLAG_LEBAR, FLAG_DATA)
    if CATATAN_HARGA:
        _paragraf(doc, "Catatan khusus perbandingan harga:", tebal=True)
        for t in CATATAN_HARGA:
            _poin(doc, t)

    # ---------------- 6. Contoh email ----------------
    _judul(doc, "6. Contoh Hasil Email")
    _paragraf(doc, "Satu email dikirim untuk setiap grup approver, berisi hanya PO yang menjadi "
                   "tanggung jawab grup tersebut. Baris yang memerlukan perhatian diberi penanda "
                   "dan latar berbeda. Setiap email disertai lampiran Excel berisi rincian yang sama.")
    _gambar(doc, png_email, 6.35,
            "Gambar 2 — Contoh tampilan email notifikasi")
    _judul(doc, "Isi email", 2)
    for t in ISI_EMAIL_POIN:
        _poin(doc, t)

    # ---------------- 7. Batasan ----------------
    _judul(doc, "7. Batasan dan Hal yang Masih Terbuka")
    _paragraf(doc, "Bagian ini disampaikan secara terbuka agar dapat menjadi bahan keputusan bersama.")
    _tabel(doc, BATASAN_JUDUL, BATASAN_LEBAR,
           [b[:2] for b in BATASAN_DATA], warna_baris=[b[2] for b in BATASAN_DATA])
    _paragraf(doc, "Tim ABAP siap menindaklanjuti butir-butir di atas sesuai arahan dan "
                   "prioritas yang ditetapkan.")

    doc.save(path_keluar)


# ================================ JALANKAN ===================================
def main():
    folder = os.path.dirname(os.path.abspath(__file__))
    png_alur = os.path.join(folder, "fs_diagram_alur.png")
    png_email = os.path.join(folder, "fs_diagram_email.png")
    keluar = os.path.join(folder, DOK["file_keluar"])

    print("1/3  Membuat diagram alur ...")
    gambar_alur(png_alur)

    print("2/3  Membuat ilustrasi email ...")
    gambar_email(png_email)

    print("3/3  Menyusun dokumen Word ...")
    buat_dokumen(png_alur, png_email, keluar)

    print("\nSelesai. Berkas dihasilkan:")
    print("   -", keluar)
    print("   -", png_alur)
    print("   -", png_email)


if __name__ == "__main__":
    main()
